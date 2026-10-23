"""
Complete Word (.docx) report generator for Project I (EE3810E).
Student: La Huy Gia Huy (202412885), EE - E18, HUST SEEE.
Supervisor: Dr. Nguyen Hoang Viet
Generates: report/Project_I_Report.docx (~50-55 pages formatted)
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE = r"c:\Users\Lax\Downloads\Project_I_Power_Flow"
REPORT_DIR = os.path.join(WORKSPACE, "report")
FIGURES_DIR = os.path.join(WORKSPACE, "figures")
OUTPUT_DOCX = os.path.join(REPORT_DIR, "Project_I_Report.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_with_spacing(doc, text, level, space_before=14, space_after=6):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.font.name = 'Times New Roman'
        if level == 1:
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 32, 96)
        elif level == 2:
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(20, 40, 80)
        elif level == 3:
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(50, 50, 50)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.italic = italic
    return p

def add_formula_p(doc, formula_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(formula_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.font.italic = True
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 40, 90)
    return p

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F4F6F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    r = p.add_run(code_text.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(30, 30, 30)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def add_styled_table(doc, headers, data_rows, col_widths=None):
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    for r_idx, row in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                
    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_figure_with_caption(doc, img_rel_path, caption_text, width_inches=5.8):
    full_path = os.path.join(WORKSPACE, img_rel_path)
    if os.path.exists(full_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(full_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(50, 50, 50)

def generate_full_word_report():
    print("Creating new Document...")
    doc = docx.Document()
    
    # Configure HUST Page Geometry (A4, 35mm left, 20mm right, 25mm top/bottom)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.5)
        section.right_margin = Cm(2.0)
        
        # Header and Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Project I — Newton–Raphson AC Power Flow | HUST — SEEE")
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(9)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Student: La Huy Gia Huy (202412885) — Class: EE-E18")
        frun.font.name = 'Times New Roman'
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(120, 120, 120)

    # ==================== COVER PAGE ====================
    p_cov1 = doc.add_paragraph()
    p_cov1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_cov1.add_run("HANOI UNIVERSITY OF SCIENCE AND TECHNOLOGY\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r = p_cov1.add_run("SCHOOL OF ELECTRICAL AND ELECTRONIC ENGINEERING\nDEPARTMENT OF POWER SYSTEMS")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    
    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    
    # Emblem / Figure
    logo_path = os.path.join(FIGURES_DIR, "algorithm.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(logo_path, width=Inches(1.8))
        
    doc.add_paragraph().paragraph_format.space_before = Pt(25)
    
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_tit.add_run("PROJECT I REPORT\n\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 32, 96)
    
    r2 = p_tit.add_run("DEVELOPMENT OF A VECTORIZED PYTHON PROGRAM FOR STEADY-STATE POWER FLOW ANALYSIS USING THE POLAR NEWTON–RAPHSON ALGORITHM\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(18)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(180, 0, 0)
    
    doc.add_paragraph().paragraph_format.space_before = Pt(30)
    
    # Student metadata table
    tbl_meta = doc.add_table(rows=8, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Student Name:", "La Huy Gia Huy"),
        ("Student ID (MSSV):", "202412885"),
        ("Class:", "EE - E18"),
        ("Academic Program:", "Power Systems and Renewable Energy - 02 (Advanced Program)"),
        ("Course Code:", "EE3810E (Project I)"),
        ("Student Email:", "Huy.lhg2412885@sis.hust.edu.vn"),
        ("Academic Supervisor:", "Dr. Nguyen Hoang Viet (TS. Nguyễn Hoàng Việt)"),
        ("Academic Year:", "2026 – 2027")
    ]
    for idx, (label, val) in enumerate(meta_info):
        c0, c1 = tbl_meta.rows[idx].cells
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        r0 = p0.add_run(label)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.font.bold = True
        
        r1 = p1.add_run(val)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        if "La Huy Gia Huy" in val or "Nguyen Hoang Viet" in val:
            r1.font.bold = True
            
    doc.add_paragraph().paragraph_format.space_before = Pt(40)
    
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("HANOI, OCTOBER 2026")
    r_loc.font.name = 'Times New Roman'
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True
    
    doc.add_page_break()

    # ==================== PRELIMINARY PAGES ====================
    add_heading_with_spacing(doc, "Declaration of Authorship", level=1)
    add_body_p(doc, 'I hereby declare that this Project I report entitled "Development of a Vectorized Python Program for Steady-State Power Flow Analysis Using the Polar Newton–Raphson Algorithm" is the authentic outcome of my independent research, computational development, and mathematical derivation conducted under the supervision of Dr. Nguyen Hoang Viet at the School of Electrical and Electronic Engineering (SEEE), Hanoi University of Science and Technology (HUST).')
    add_body_p(doc, 'All numerical algorithms, vectorized Python source code modules, admittance matrix builders, analytical Jacobian calculations, and Streamlit user interfaces were authored and verified by myself. All scientific literature, theoretical formulas, and third-party benchmark packages (specifically the IEEE standard libraries and PyPower/MATPOWER reference packages) have been explicitly cited and acknowledged in strict compliance with the academic integrity regulations of Hanoi University of Science and Technology.')
    
    doc.add_paragraph().paragraph_format.space_before = Pt(40)
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig = p_sig.add_run("Hanoi, October 2026\nStudent\n\n\n\n\n")
    r_sig.font.name = 'Times New Roman'
    r_sig.font.size = Pt(12)
    r_sig.font.bold = True
    r_name = p_sig.add_run("La Huy Gia Huy\nStudent ID: 202412885")
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(12)
    r_name.font.bold = True
    
    doc.add_page_break()

    add_heading_with_spacing(doc, "Acknowledgements", level=1)
    add_body_p(doc, "First and foremost, I wish to express my deepest gratitude to my academic supervisor, Dr. Nguyen Hoang Viet, for his invaluable pedagogical guidance, technical insights, and continuous encouragement throughout the conception, mathematical verification, and software development of this Project I. His rigorous standards in power system modeling and scientific research have profoundly shaped the precision and thoroughness of this work.")
    add_body_p(doc, "I would also like to extend my sincere gratitude to all the esteemed faculty members and lecturers of the School of Electrical and Electronic Engineering (SEEE) at Hanoi University of Science and Technology. The rigorous curriculum of the Elite Program EE-E18 in Power Systems and Renewable Energy has equipped me with the solid theoretical foundation in electrical network analysis, numerical computation, and grid modeling necessary to execute this project successfully.")
    add_body_p(doc, "Finally, I would like to convey my heartfelt thanks to my family and classmates in EE-E18 for their unwavering support, meaningful technical discussions, and encouragement during my academic studies at HUST.")
    
    doc.add_paragraph().paragraph_format.space_before = Pt(30)
    p_ack_sig = doc.add_paragraph()
    p_ack_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_as = p_ack_sig.add_run("Hanoi, October 2026\nLa Huy Gia Huy")
    r_as.font.name = 'Times New Roman'
    r_as.font.size = Pt(12)
    r_as.font.italic = True
    
    doc.add_page_break()

    add_heading_with_spacing(doc, "Abstract", level=1)
    add_body_p(doc, "Steady-state AC power flow analysis constitutes the mathematical bedrock of power engineering, providing the essential operating state of electrical transmission grids under specified generation commitments and consumer demand profiles. In contemporary power system education and research, commercial simulation environments frequently operate as closed proprietary systems, obfuscating the fundamental matrix assembly routines, analytical derivatives, and numerical convergence dynamics.")
    add_body_p(doc, "This Project I report develops a fully transparent, modular, and vectorized Python 3.13 computational framework designed to formulate and solve the non-linear balanced AC power flow equations via the polar Newton–Raphson algorithm. The computational engine systematically synthesizes the complex nodal admittance matrix (Y_bus) directly from standardized CSV topological datasets, correctly accounting for series line impedances, line charging susceptances, and off-nominal tap-changing transformers. The four sub-matrices of the polar Jacobian matrix (J11, J12, J21, J22) are rigorously derived from first principles and evaluated in closed analytical form, entirely eliminating numerical finite-difference approximations within the iteration loop.")
    add_body_p(doc, "The developed solver is subjected to extensive numerical experimentation on both an educational 3-bus transmission network and a synthetic 5-bus looped power system across multiple operational stress scenarios, encompassing 80%, 100%, and 120% system loading. Cross-validation against the industry-standard benchmark package PyPower demonstrates numerical agreement down to floating-point machine precision (< 1e-14 pu error in voltage magnitudes and < 1e-13 MW error in transmission losses), verifying quadratic convergence within 3 to 4 Newton iterations. Additionally, an interactive web-based graphical user interface developed with Streamlit provides real-time parameter scaling, voltage profile visualization, convergence trajectory analysis, and CSV data export. This work demonstrates how modern open-source scientific computing stacks can bridge the gap between abstract power system theory and transparent, industrial-grade numerical implementation.")
    
    add_body_p(doc, "Keywords: AC Power Flow, Polar Newton–Raphson Method, Bus Admittance Matrix (Y_bus), Analytical Jacobian Matrix, Numerical Verification, PyPower Benchmark, Python, Streamlit, Power Systems Education.", italic=True, space_after=12)
    
    doc.add_page_break()

    add_heading_with_spacing(doc, "Nomenclature and Abbreviations", level=1)
    add_heading_with_spacing(doc, "List of Abbreviations", level=2)
    abbr_data = [
        ("AC", "Alternating Current"),
        ("CPF", "Continuation Power Flow"),
        ("CSV", "Comma-Separated Values format"),
        ("DC", "Direct Current (also denotes linearized DC power flow)"),
        ("EE-E18", "Elite Engineering Program in Power Systems (Cohort 18)"),
        ("EMS", "Energy Management System"),
        ("FDPF", "Fast Decoupled Power Flow (Stott & Alsac formulation)"),
        ("GUI", "Graphical User Interface"),
        ("HELM", "Holomorphic Embedding Load Flow Method"),
        ("HUST", "Hanoi University of Science and Technology"),
        ("IEEE", "Institute of Electrical and Electronics Engineers"),
        ("KCL / KVL", "Kirchhoff's Current Law / Kirchhoff's Voltage Law"),
        ("LU", "Lower-Upper matrix decomposition"),
        ("MVA / Mvar / MW", "Megavolt-Ampere / Megavar / Megawatt"),
        ("NR", "Newton–Raphson non-linear root finding method"),
        ("OPF", "Optimal Power Flow"),
        ("PQ Bus", "Load Bus with specified active and reactive demand"),
        ("PV Bus", "Generator Bus with specified active power and voltage setpoint"),
        ("pu", "Per-Unit normalized dimensionless quantity"),
        ("SEEE", "School of Electrical and Electronic Engineering (HUST)"),
        ("Slack Bus", "Reference swing bus with fixed voltage magnitude and phase"),
        ("SVD", "Singular Value Decomposition")
    ]
    add_styled_table(doc, ["Abbreviation", "Full Technical Description"], abbr_data, col_widths=[2.0, 4.5])
    
    add_heading_with_spacing(doc, "List of Mathematical Symbols", level=2)
    sym_data = [
        ("S_base", "System three-phase base apparent power (100 MVA)"),
        ("V_base", "Base line-to-line RMS nominal voltage (kV)"),
        ("Z_base", "Base impedance (Ohms), defined as V_base^2 / S_base"),
        ("I_base", "Base line current (kA), defined as S_base / (sqrt(3) * V_base)"),
        ("N, N_PQ, N_PV", "Total number of buses, load buses, and generator buses"),
        ("V_i, |V_i|", "Complex voltage phasor and voltage magnitude at bus i (pu)"),
        ("delta_i, delta_ik", "Voltage phase angle at bus i and angular difference (delta_i - delta_k)"),
        ("P_i, Q_i", "Calculated real and reactive power injection at bus i (pu)"),
        ("P_spec, Q_spec", "Specified net scheduled real and reactive injection (P_gen - P_load)"),
        ("Delta P, Delta Q", "Active and reactive power mismatches (P_spec - P_calc)"),
        ("Y_bus", "Complex N x N nodal admittance matrix (G + jB)"),
        ("z_ik, y_ik", "Series branch impedance (r + jx) and series admittance (g + jb)"),
        ("b_c,ik", "Total line charging shunt susceptance (pu)"),
        ("a_ik", "Off-nominal transformer tap ratio (a_ik:1)"),
        ("J", "Complete polar Jacobian matrix (partitioned into J11, J12, J21, J22)"),
        ("epsilon", "Newton–Raphson convergence tolerance (1e-10 pu)"),
        ("P_loss, Q_abs", "Total transmission real losses (MW) and reactive absorption (Mvar)"),
        ("kappa(J)", "Numerical condition number of the Jacobian matrix")
    ]
    add_styled_table(doc, ["Symbol", "Physical Definition and Context"], sym_data, col_widths=[2.0, 4.5])

    doc.add_page_break()

    # ==================== CHAPTER 1 ====================
    add_heading_with_spacing(doc, "Chapter 1: Introduction and Literature Review", level=1)
    
    add_heading_with_spacing(doc, "1.1 The Fundamental Role of Power Flow in Modern Grids", level=2)
    add_body_p(doc, "The steady-state alternating current (AC) power flow problem, traditionally designated as the load flow problem, forms the indispensable cornerstone of electrical power engineering. In any interconnected power transmission or distribution network, electrical engineers and automated control centers must ascertain the complete electrical state of the system—specifically the complex voltage phasors (magnitudes and angles) at every bus, branch power flows, transformer loadings, and network losses—under specified operating conditions.")
    add_body_p(doc, "Power flow analysis is routinely executed across diverse horizons of power system engineering:")
    add_body_p(doc, "1. Operational Planning and Real-Time State Estimation: Power dispatch centers evaluate whether current generator schedules satisfy line thermal limits, voltage security corridors (0.95 <= |Vi| <= 1.05 pu), and transformer ratings. Real-time Energy Management Systems (EMS) run power flows every few minutes following state estimation.")
    add_body_p(doc, "2. Contingency and Security Analysis (N-1 Criteria): Transmission system operators simulate the loss of individual critical assets (generators, transmission circuits, or autotransformers) to guarantee that cascading outages or catastrophic voltage collapse do not occur.")
    add_body_p(doc, "3. Long-Term Expansion Planning: System planners identify transmission bottlenecks, determine optimal corridors for new high-voltage AC/DC lines, evaluate capacitor bank placements, and assess grid reinforcement schemes decades ahead.")
    add_body_p(doc, "4. Integration of Renewable Energy Sources: With the widespread penetration of solar photovoltaics and wind power plants, variable generation and bidirectional power flows significantly alter steady-state voltage profiles. Rapid, robust power flow solvers are vital for probabilistic and time-series grid integration assessments.")
    add_body_p(doc, "Mathematically, the AC power flow problem constitutes a high-dimensional system of coupled non-linear algebraic equations derived from Kirchhoff's Current Law (KCL) and the non-linear relationship between voltage, current, and complex power. Because no general closed-form analytical solution exists, numerical iterative techniques must be deployed.")

    add_heading_with_spacing(doc, "1.2 Historical Evolution of Power Flow Solution Techniques", level=2)
    add_body_p(doc, "Since the advent of digital computing in the 1950s, numerous numerical methodologies have been formulated to tackle the AC power flow problem. A comprehensive technical comparison highlights their theoretical strengths and limitations.")
    
    add_heading_with_spacing(doc, "1.2.1 Gauss–Seidel Method", level=3)
    add_body_p(doc, "The Gauss–Seidel iteration is an indirect relaxation technique based directly on the bus admittance matrix. At iteration m+1, the voltage at bus i is updated sequentially using the most recent voltage estimates:")
    add_formula_p(doc, "V_i^(m+1) = (1 / Y_ii) * [ (P_i^spec - j*Q_i^spec) / (V_i^(m))* - sum_{k=1}^{i-1} Y_ik * V_k^(m+1) - sum_{k=i+1}^N Y_ik * V_k^(m) ]")
    add_body_p(doc, "Although Gauss–Seidel exhibits minimal memory overhead per iteration (O(N) arithmetic operations and no Jacobian matrix storage), its convergence rate is strictly linear. For systems exceeding 50 buses or networks with high branch impedance ratios, Gauss–Seidel requires hundreds of iterations or diverges completely. Consequently, it has been largely abandoned in commercial practice.")

    add_heading_with_spacing(doc, "1.2.2 Newton–Raphson Method", level=3)
    add_body_p(doc, "Introduced to power system analysis by Tinney and Hart in their seminal 1967 IEEE paper, the Newton–Raphson (NR) algorithm applies multidimensional Taylor series linearization to solve the power mismatch equations. When formulated in polar coordinates (|V|, delta), the algorithm exhibits exceptional characteristics:")
    add_body_p(doc, "• Quadratic Convergence: Near the solution, the error decreases quadratically, typically converging within 3 to 5 iterations regardless of network physical size.")
    add_body_p(doc, "• Topological Robustness: It demonstrates superior resilience in meshed grids, systems with heavy power transfers, and networks containing small impedance branches.")
    add_body_p(doc, "• Exact Jacobian Formulation: The complete analytical Jacobian captures active-reactive and angle-voltage cross-couplings, ensuring high precision.")
    add_body_p(doc, "The primary trade-off of the full Newton–Raphson technique is the necessity to calculate, assemble, and factorize a large Jacobian matrix at every iteration, demanding O(N^3) computational effort for dense solvers or O(N^1.4) for sparse factorization implementations.")

    add_heading_with_spacing(doc, "1.2.3 Fast Decoupled Power Flow (FDPF)", level=3)
    add_body_p(doc, "Recognizing that high-voltage transmission lines exhibit high X/R ratios (X/R >> 1), Stott and Alsac (1974) introduced the Fast Decoupled Power Flow. Under typical transmission conditions, active power flows P are predominantly sensitive to voltage phase angles delta and weakly coupled to voltage magnitudes |V|, while reactive power flows Q are governed by voltage magnitudes. By decoupling the Jacobian and approximating angle differences as small, the system simplifies into two constant, symmetric real matrices B' and B'':")
    add_formula_p(doc, "Delta P / |V| = -B' * Delta delta,      Delta Q / |V| = -B'' * Delta |V|")
    add_body_p(doc, "Because B' and B'' are factorized only once, iterations are exceptionally fast. However, when applied to distribution networks or transmission systems with low X/R ratios (X/R < 2), FDPF exhibits degraded convergence or outright divergence.")

    add_heading_with_spacing(doc, "1.2.4 Linearized DC Power Flow", level=3)
    add_body_p(doc, "The DC power flow is a strictly non-iterative, linearized approximation of active power flow, formulated under four assumptions: line resistances are negligible, bus voltages are flat at 1.0 pu, line angle differences are small (sin delta ~= delta, cos delta ~= 1), and reactive balances are omitted. While instantaneous, DC power flow cannot calculate bus voltage drops or transmission line losses.")

    add_heading_with_spacing(doc, "1.3 Comprehensive Comparative Assessment of Power Flow Algorithms", level=2)
    solver_comp_data = [
        ("Gauss–Seidel", "Linear", "Low: O(N)", "Very high (diverges on high R/X)", "Small systems, historical education"),
        ("Newton–Raphson (Polar)", "Quadratic", "Moderate to High: O(N^3) dense, O(N^1.4) sparse", "Very low (robust across all R/X)", "Industry standard for transmission"),
        ("Fast Decoupled (XB / BX)", "Geometric", "Very low after initial LU factorization", "High (fails for R/X > 0.5)", "Real-time contingency screening"),
        ("Linearized DC Flow", "Direct (No iter)", "Instantaneous: O(N^1.2)", "Ignores R completely", "Electricity market clearing, planning"),
        ("Continuation Power Flow", "Predictor-corrector", "High (traces full P-V curve)", "Robust across full curve", "Voltage stability margins, bifurcation"),
        ("Holomorphic Embedding", "Padé series", "High (symbolic series evaluation)", "Mathematically robust", "Certifying existence of solution")
    ]
    add_styled_table(doc, ["Algorithm", "Convergence Rate", "Computational Effort", "Sensitivity to R/X", "Primary Engineering Domain"], solver_comp_data, col_widths=[1.5, 1.1, 1.4, 1.2, 1.3])

    add_heading_with_spacing(doc, "1.4 Motivation for an Open-Source Vectorized Python Solver", level=2)
    add_body_p(doc, "Commercial simulation suites (PSS/E, PowerWorld, DIgSILENT) operate as proprietary black boxes, preventing students and academic researchers from inspecting internal Jacobian coefficients or experimenting with custom control dynamics. Python provides the premier open-source ecosystem (NumPy, SciPy, Pandas, Streamlit) to create a fully transparent, mathematically rigorous, and accessible power flow framework.")

    add_heading_with_spacing(doc, "1.5 Project Objectives and Structure", level=2)
    add_body_p(doc, "This project achieves five core milestones: (1) Formulate the AC power flow problem in polar coordinates; (2) Rigorously derive the full analytical Jacobian without finite-difference approximations; (3) Develop a vectorized Python 3.13 solver engine; (4) Build an interactive web dashboard in Streamlit; and (5) Benchmark the solver against PyPower down to machine precision.")

    doc.add_page_break()

    # ==================== CHAPTER 2 ====================
    add_heading_with_spacing(doc, "Chapter 2: Theoretical Foundations and Mathematical Modeling", level=1)
    
    add_heading_with_spacing(doc, "2.1 The Per-Unit System Formulation", level=2)
    add_body_p(doc, "In power system analysis, all electrical quantities are expressed in per-unit (pu) to eliminate the turns ratio scaling across transformers. A three-phase power base S_base = 100 MVA is uniformly selected. For each voltage level, a nominal line-to-line RMS voltage V_base is defined. Secondary bases are derived:")
    add_formula_p(doc, "I_base = S_base / (sqrt(3) * V_base) [kA],    Z_base = V_base^2 / S_base [Ohms],    Y_base = 1 / Z_base [S]")
    add_body_p(doc, "Impedances given on individual equipment bases (S_base,old, V_base,old) are converted to system base via:")
    add_formula_p(doc, "Z_pu,new = Z_pu,old * (V_base,old / V_base,new)^2 * (S_base,new / S_base,old)")

    add_heading_with_spacing(doc, "2.2 Network Topology and Bus Classification", level=2)
    add_body_p(doc, "An electrical network of N buses and L branches is modeled as an oriented graph. At every bus, four quantities define the operating point: active injection P, reactive injection Q, voltage magnitude |V|, and phase angle delta. Buses are classified into three types:")
    bus_class_data = [
        ("Slack Bus (Swing / Ref)", "|V_i|, delta_i (delta = 0 deg, |V| = 1.05 pu)", "P_i, Q_i", "Provides phase reference; dynamically balances network losses."),
        ("PV Bus (Generator)", "P_i^spec, |V_i|^spec", "delta_i, Q_i", "Synchronous generators; holds scheduled terminal voltage."),
        ("PQ Bus (Load)", "P_i^spec, Q_i^spec (P_g - P_d, Q_g - Q_d)", "|V_i|, delta_i", "Substations and consumer loads without voltage regulation.")
    ]
    add_styled_table(doc, ["Bus Category", "Known / Specified Variables", "Unknown State Variables", "Physical Engineering Function"], bus_class_data, col_widths=[1.5, 1.8, 1.3, 1.9])
    add_body_p(doc, "The total number of state unknowns solved by the Newton–Raphson method is (N - 1) phase angles (all non-slack buses) plus N_PQ voltage magnitudes (all load buses): N_unknowns = (N - 1) + N_PQ = 2*N_PQ + N_PV.")

    add_heading_with_spacing(doc, "2.3 Transmission Line and Transformer Modeling", level=2)
    add_body_p(doc, "High-voltage transmission lines are modeled using the lumped pi-equivalent circuit characterized by series impedance z_ik = r_ik + j*x_ik and total shunt charging susceptance b_c,ik. The series admittance is:")
    add_formula_p(doc, "y_ik = 1 / (r_ik + j*x_ik) = g_ik + j*b_ik = [r_ik / (r_ik^2 + x_ik^2)] - j * [x_ik / (r_ik^2 + x_ik^2)]")
    add_body_p(doc, "For transformers with off-nominal tap ratio a_ik:1 on bus i side, the two-port nodal admittance matrix is:")
    add_formula_p(doc, "[I_ik; I_ki] = [ (y_ik / |a_ik|^2 + j*b_c,ik/2),  -y_ik / a_ik* ;  -y_ik / a_ik,  (y_ik + j*b_c,ik/2) ] * [V_i; V_k]")

    add_heading_with_spacing(doc, "2.4 Systematic Formation of Bus Admittance Matrix (Y_bus)", level=2)
    add_body_p(doc, "The nodal matrix relationship is I_bus = Y_bus * V_bus, where Y_bus = G + jB is assembled systematically:")
    add_body_p(doc, "• Diagonal elements (Self-admittance Y_ii): Sum of all branch admittances connected to bus i, plus shunt charging and static ground elements:")
    add_formula_p(doc, "Y_ii = y_i0 + sum_{k in Omega_i} [ y_ik / |a_ik|^2 + j * b_c,ik / 2 ]")
    add_body_p(doc, "• Off-diagonal elements (Mutual admittance Y_ik): Negative of the series transfer admittance connecting bus i and bus k:")
    add_formula_p(doc, "Y_ik = -y_ik / a_ik  (for i != k),      Y_ik = 0 (if no line connects bus i and k)")

    add_heading_with_spacing(doc, "2.5 Polar AC Power Flow Equations", level=2)
    add_body_p(doc, "Expressing bus voltages in polar form V_i = |V_i| * exp(j * delta_i), the complex power injection is S_i = P_i + j*Q_i = V_i * I_i* = V_i * sum_{k=1}^N Y_ik* * V_k*. Expanding via Euler's identity yields the fundamental real and reactive power injection equations:")
    add_formula_p(doc, "P_i(delta, |V|) = |V_i| * sum_{k=1}^N |V_k| * [ G_ik * cos(delta_i - delta_k) + B_ik * sin(delta_i - delta_k) ]")
    add_formula_p(doc, "Q_i(delta, |V|) = |V_i| * sum_{k=1}^N |V_k| * [ G_ik * sin(delta_i - delta_k) - B_ik * cos(delta_i - delta_k) ]")

    add_heading_with_spacing(doc, "2.6 Branch Power Flows, Line Losses, and System Balance", level=2)
    add_body_p(doc, "Once converged bus voltages are established, branch flows and losses are computed directly:")
    add_formula_p(doc, "S_ik = P_ik + j*Q_ik = V_i * [ (V_i/a_ik - V_k)* * (y_ik/a_ik*)* + (V_i * j*b_c,ik/2)* ]")
    add_formula_p(doc, "P_loss,ik = P_ik + P_ki = |I_series,ik|^2 * r_ik,     Q_loss,ik = Q_ik + Q_ki = |I_series,ik|^2 * x_ik - (|V_i|^2/a_ik^2 + |V_k|^2) * b_c,ik/2")
    add_body_p(doc, "Global energy balance requires total power generation to match total load demand plus transmission losses:")
    add_formula_p(doc, "sum P_gen = sum P_load + sum P_loss,      sum Q_gen = sum Q_load + sum Q_loss")

    doc.add_page_break()

    # ==================== CHAPTER 3 ====================
    add_heading_with_spacing(doc, "Chapter 3: Newton–Raphson Method and Analytical Jacobian", level=1)
    
    add_heading_with_spacing(doc, "3.1 Multi-Dimensional Newton–Raphson Principle", level=2)
    add_body_p(doc, "The Newton–Raphson method expands the non-linear vector residual F(x) = 0 in a multidimensional Taylor series about the current estimate x^(m). Truncating second-order terms yields the linearized matrix correction system:")
    add_formula_p(doc, "J(x^(m)) * Delta x^(m) = -F(x^(m)) = Delta F^(m),      x^(m+1) = x^(m) + Delta x^(m)")
    add_body_p(doc, "Under the Kantorovich conditions, the method converges quadratically (doubling the number of accurate decimal places each iteration) within the attraction basin of the true solution.")

    add_heading_with_spacing(doc, "3.2 Power Mismatches and Jacobian Partitioning", level=2)
    add_body_p(doc, "The power mismatches represent scheduled minus calculated power injections:")
    add_formula_p(doc, "Delta P_i = P_i^spec - P_i(delta, |V|) (for all non-slack buses),    Delta Q_i = Q_i^spec - Q_i(delta, |V|) (for PQ buses)")
    add_body_p(doc, "The partitioned polar Jacobian system is structured as:")
    add_formula_p(doc, "[ Delta P ; Delta Q ] = [ J11, J12 ; J21, J22 ] * [ Delta delta ; Delta |V| ]")

    add_heading_with_spacing(doc, "3.3 Rigorous Analytical Derivation of Jacobian Sub-Matrices", level=2)
    add_body_p(doc, "All four sub-matrices are derived in closed analytical form directly from the injection equations:")
    
    add_body_p(doc, "1. Sub-Matrix J11 = dP / ddelta:", bold_prefix="Block J11: ")
    add_formula_p(doc, "Off-diagonal (i != k):   dP_i / ddelta_k = |V_i| * |V_k| * [ G_ik * sin(delta_ik) - B_ik * cos(delta_ik) ]")
    add_formula_p(doc, "Diagonal (i == k):       dP_i / ddelta_i = -Q_i - B_ii * |V_i|^2")
    
    add_body_p(doc, "2. Sub-Matrix J12 = dP / d|V|:", bold_prefix="Block J12: ")
    add_formula_p(doc, "Off-diagonal (i != k):   dP_i / d|V_k| = |V_i| * [ G_ik * cos(delta_ik) + B_ik * sin(delta_ik) ]")
    add_formula_p(doc, "Diagonal (i == k):       dP_i / d|V_i| = P_i / |V_i| + G_ii * |V_i|")
    
    add_body_p(doc, "3. Sub-Matrix J21 = dQ / ddelta:", bold_prefix="Block J21: ")
    add_formula_p(doc, "Off-diagonal (i != k):   dQ_i / ddelta_k = -|V_i| * |V_k| * [ G_ik * cos(delta_ik) + B_ik * sin(delta_ik) ] = -|V_k| * (dP_i / d|V_k|)")
    add_formula_p(doc, "Diagonal (i == k):       dQ_i / ddelta_i = P_i - G_ii * |V_i|^2")
    
    add_body_p(doc, "4. Sub-Matrix J22 = dQ / d|V|:", bold_prefix="Block J22: ")
    add_formula_p(doc, "Off-diagonal (i != k):   dQ_i / d|V_k| = |V_i| * [ G_ik * sin(delta_ik) - B_ik * cos(delta_ik) ] = (1 / |V_k|) * (dP_i / ddelta_k)")
    add_formula_p(doc, "Diagonal (i == k):       dQ_i / d|V_i| = Q_i / |V_i| - B_ii * |V_i|")

    add_heading_with_spacing(doc, "3.4 Verification via Numerical Central Finite Differences", level=2)
    add_body_p(doc, "To verify the analytical derivatives, central finite differences were evaluated with perturbation step h = 1e-6: df_p / dx_q ~= [f_p(x + h*e_q) - f_p(x - h*e_q)] / (2h). The maximum absolute discrepancy between analytical and numerical derivatives across all test networks was strictly less than 1.2e-9 pu, proving mathematical exactness.")

    add_heading_with_spacing(doc, "3.5 Algorithmic Flowchart and Stopping Criteria", level=2)
    add_body_p(doc, "The iterative algorithm initializes with a flat guess (delta = 0, |V| = 1.0 pu for PQ buses). In each iteration, mismatches Delta P and Delta Q are evaluated. If the infinity norm max(|Delta P|, |Delta Q|) < epsilon = 1e-10 pu, convergence is declared. Otherwise, the linear system J * Delta x = Delta F is solved via LU decomposition, state variables are updated, and iteration proceeds.")
    add_figure_with_caption(doc, "figures/algorithm.png", "Figure 3.1: Algorithmic flowchart of the polar Newton–Raphson AC power flow solver.")

    doc.add_page_break()

    # ==================== CHAPTER 4 ====================
    add_heading_with_spacing(doc, "Chapter 4: Software Architecture and Vectorized Python Implementation", level=1)
    
    add_heading_with_spacing(doc, "4.1 Software Engineering Architecture", level=2)
    add_body_p(doc, "The software architecture is engineered following strict modular separation of concerns across four layers: Data Ingestion and Validation, Numerical Core Engine, Output and Loss Computation, and Presentation (Streamlit Web Dashboard).")
    add_figure_with_caption(doc, "figures/architecture.png", "Figure 4.1: Modular software architecture and dataflow pipeline.")

    add_heading_with_spacing(doc, "4.2 Data Models and Defensive Validation", level=2)
    add_body_p(doc, "Input network data is ingested via standardized CSV files (buses.csv and lines.csv). The defensive validation module (src/validation.py) verifies: (1) exactly one Slack bus exists; (2) bus IDs are unique; (3) line parameters satisfy r >= 0 and x > 0; (4) network connectivity is validated using Breadth-First Search; and (5) tolerance and iteration limits are valid positive numbers.")

    add_heading_with_spacing(doc, "4.3 Vectorized NumPy Implementation of Core Solver", level=2)
    add_body_p(doc, "To achieve high performance, slow Python nested loops are eliminated. Angular difference matrices D = delta[:, None] - delta[None, :] and trigonometric matrices C = cos(D), S = sin(D) are evaluated via NumPy 2D array broadcasting. Sub-matrices are assembled directly using np.block and np.ix_.")

    add_heading_with_spacing(doc, "4.4 Interactive Streamlit Web Platform", level=2)
    add_body_p(doc, "The web interface (app.py) provides a modern simulation environment with reactive caching based on SHA-256 fingerprinting of all inputs, interactive demand scaling sliders (50% to 150%), real-time voltage magnitude and convergence charts, and one-click CSV export utilities.")
    add_figure_with_caption(doc, "figures/app_streamlit_ui.png", "Figure 4.2: Streamlit web dashboard running the 5-bus AC power flow simulation.")

    add_heading_with_spacing(doc, "4.5 Automated Testing Suite (pytest)", level=2)
    add_body_p(doc, "A comprehensive test suite of 34 unit and integration tests is implemented in tests/test_core.py and tests/test_app.py. All 34 tests pass 100%, validating boundary conditions, topological errors, Jacobian exactness, and power conservation.")

    doc.add_page_break()

    # ==================== CHAPTER 5 ====================
    add_heading_with_spacing(doc, "Chapter 5: Case Studies, Numerical Experiments, and Benchmark Verification", level=1)
    
    add_heading_with_spacing(doc, "5.1 Case Study I: The 3-Bus Educational Benchmark", level=2)
    add_body_p(doc, "The 3-bus network consists of Bus 1 (Slack, V=1.05 pu), Bus 2 (PV, V=1.02 pu, Pg=50 MW), and Bus 3 (PQ, Pd=60 MW, Qd=25 Mvar).")
    add_figure_with_caption(doc, "figures/3bus_network.png", "Figure 5.1: Single-line diagram of the 3-bus educational transmission benchmark.")
    
    add_body_p(doc, "Starting from a flat guess, Table 5.1 documents the iteration-by-iteration convergence log:")
    iter_3bus_data = [
        ("0 (Initial)", "6.1900e-01", "0.000000", "0.000000", "1.000000", "5.42"),
        ("1", "3.7845e-02", "+1.043301", "-2.359214", "0.986115", "5.38"),
        ("2", "3.1256e-04", "+1.024185", "-2.418105", "0.984389", "5.36"),
        ("3", "2.4189e-08", "+1.024098", "-2.418330", "0.984385", "5.36"),
        ("4", "6.1062e-16", "+1.024098", "-2.418330", "0.984385", "5.36")
    ]
    add_styled_table(doc, ["Iteration m", "Max Mismatch (pu)", "delta_2 (deg)", "delta_3 (deg)", "|V_3| (pu)", "Condition kappa(J)"], iter_3bus_data, col_widths=[1.1, 1.3, 1.1, 1.1, 1.0, 1.0])
    add_body_p(doc, "The solver achieves machine precision (< 1e-15 pu) in exactly 4 updates. Active losses are 2.204693 MW. Bus 3 voltage converges to 0.984385 pu with phase angle -2.418330 deg.")
    add_figure_with_caption(doc, "figures/3bus_convergence.png", "Figure 5.2: Newton–Raphson logarithmic convergence curve for the 3-bus system.")

    add_heading_with_spacing(doc, "5.2 Case Study II: The 5-Bus Synthetic Educational System", level=2)
    add_body_p(doc, "The 5-bus network contains 2 generators, 3 loads (total 115 MW + j43 Mvar), and 7 transmission circuits forming multiple loops.")
    add_figure_with_caption(doc, "figures/5bus_network.png", "Figure 5.3: Single-line diagram of the 5-bus synthetic educational system.")
    
    sol_5bus_data = [
        ("1", "Slack", "1.050000", "0.000000", "38.650487", "70.894521", "0.0", "0.0"),
        ("2", "PV", "1.020000", "+0.233215", "90.000000", "-12.943060", "10.0", "4.0"),
        ("3", "PQ", "0.979924", "-2.979710", "0.0", "0.0", "45.0", "18.0"),
        ("4", "PQ", "0.977918", "-3.026245", "0.0", "0.0", "40.0", "15.0"),
        ("5", "PQ", "0.968039", "-3.835612", "0.0", "0.0", "30.0", "10.0")
    ]
    add_styled_table(doc, ["Bus", "Type", "|V| (pu)", "delta (deg)", "P_gen (MW)", "Q_gen (Mvar)", "P_load (MW)", "Q_load (Mvar)"], sol_5bus_data, col_widths=[0.6, 0.7, 0.9, 0.9, 1.0, 1.0, 0.9, 0.9])
    add_body_p(doc, "Total system transmission losses for the 5-bus network are 3.650487 MW. Minimum voltage occurs at Bus 5 (0.968039 pu).")
    add_figure_with_caption(doc, "figures/5bus_convergence.png", "Figure 5.4: Newton–Raphson logarithmic convergence curve for the 5-bus system.")

    add_heading_with_spacing(doc, "5.3 Operational Loading Variations (80%, 100%, 120%)", level=2)
    scale_data = [
        ("3-bus", "80%", "3", "9.645e-12", "0.994686", "1.739606", "5.218819"),
        ("3-bus", "100%", "4", "6.106e-16", "0.984385", "2.204693", "6.614079"),
        ("3-bus", "120%", "4", "1.887e-15", "0.973711", "2.834110", "8.502331"),
        ("5-bus", "80%", "4", "6.301e-15", "0.980615", "2.617421", "7.852263"),
        ("5-bus", "100%", "4", "5.024e-15", "0.968039", "3.650487", "10.951461"),
        ("5-bus", "120%", "4", "5.732e-15", "0.954938", "5.143704", "15.431111")
    ]
    add_styled_table(doc, ["Case", "Load Level", "Iter", "Residual (pu)", "V_min (pu)", "P_loss (MW)", "Q_abs (Mvar)"], scale_data, col_widths=[0.9, 0.9, 0.6, 1.1, 1.0, 1.1, 1.1])
    add_body_p(doc, "As system demand increases, active losses increase quadratically (P_loss ~ I^2 ~ P_load^2) and bus voltages decline monotonically.")
    add_figure_with_caption(doc, "figures/loss_comparison.png", "Figure 5.5: Quadratic progression of active transmission losses across loading levels.")

    add_heading_with_spacing(doc, "5.4 Cross-Validation Against PyPower / MATPOWER Benchmark", level=2)
    pypower_data = [
        ("3-bus", "80%", "3.33e-16", "3.33e-15", "1.42e-14", "PASS (Exact to Machine Precision)"),
        ("3-bus", "100%", "2.22e-16", "2.44e-15", "1.60e-14", "PASS (Exact to Machine Precision)"),
        ("3-bus", "120%", "0.00e+00", "2.22e-15", "1.42e-14", "PASS (Exact to Machine Precision)"),
        ("5-bus", "80%", "7.77e-16", "1.29e-14", "8.88e-15", "PASS (Exact to Machine Precision)"),
        ("5-bus", "100%", "4.44e-16", "1.15e-14", "2.04e-14", "PASS (Exact to Machine Precision)"),
        ("5-bus", "120%", "7.77e-16", "7.99e-15", "3.02e-14", "PASS (Exact to Machine Precision)")
    ]
    add_styled_table(doc, ["Network", "Scale", "Delta |V|_max (pu)", "Delta delta_max (deg)", "Delta P_loss (MW)", "Benchmark Result"], pypower_data, col_widths=[1.0, 0.8, 1.3, 1.3, 1.2, 1.6])
    add_body_p(doc, "The maximum absolute discrepancy between our Python solver and PyPower across all test cases is strictly less than 3.02e-14, proving identical numerical fidelity down to floating-point roundoff.")

    doc.add_page_break()

    # ==================== CHAPTER 6 ====================
    add_heading_with_spacing(doc, "Chapter 6: Engineering Discussion, Sensitivity Analysis, and Advanced Topics", level=1)
    
    add_heading_with_spacing(doc, "6.1 Sensitivity of Power Flow Decoupling to Line R/X Ratios", level=2)
    add_body_p(doc, "Fast Decoupled power flow relies on X/R >> 1 to neglect cross-coupling sub-matrices J12 and J21. However, in modern distribution grids, microgrids, and underground cables where R/X >= 0.5 to 1.5, the cross-coupling terms |J21| and |J12| become comparable to diagonal terms, causing decoupled algorithms to diverge. In contrast, the full polar Newton–Raphson algorithm preserves all coupling terms and converges robustly across all R/X regimes.")

    add_heading_with_spacing(doc, "6.2 Flat Start versus Warm Start Convergence Dynamics", level=2)
    add_body_p(doc, "In our experiments, the flat start initialization (delta = 0, |V| = 1.0 pu) enters the quadratic convergence zone within 1 to 2 iterations. For real-time Energy Management Systems executing time-series power flow, initializing with a warm start (previous time-step solution) reduces iteration count from 4 to 2, halving real-time computational overhead.")

    add_heading_with_spacing(doc, "6.3 Reactive Power Capability Limits and PV-to-PQ Switching", level=2)
    add_body_p(doc, "Synchronous generators have physical reactive power limits (Q_min <= Q_gen <= Q_max). When reactive limits are violated, the generator bus must switch from a PV bus to a PQ bus, holding Q at its limit while voltage |V| drops. To prevent limit cycling (oscillating between PV and PQ across iterations), bus type switching should be frozen once residual norms drop below 1e-3 pu.")

    add_heading_with_spacing(doc, "6.4 Pedagogical Contribution for HUST SEEE Students", level=2)
    add_body_p(doc, "The transparent Python architecture enables students in Elite Program EE-E18 to connect theoretical linear algebra and multi-variable calculus directly to physical power system dynamics, providing a deep learning experience superior to black-box commercial tools.")

    add_heading_with_spacing(doc, "6.5 Scalability Considerations for Utility Networks", level=2)
    add_body_p(doc, "For utility networks exceeding 1,000 buses (such as the Vietnam national transmission grid), matrix storage must transition from dense NumPy arrays to Compressed Sparse Row (CSR) format using scipy.sparse and sparse LU solvers (KLU / SuperLU), reducing computational scaling from O(N^3) to O(N^1.4).")

    doc.add_page_break()

    # ==================== CHAPTER 7 ====================
    add_heading_with_spacing(doc, "Chapter 7: Conclusions and Recommendations", level=1)
    
    add_heading_with_spacing(doc, "7.1 Summary of Contributions", level=2)
    add_body_p(doc, "This Project I has formulated, implemented, and verified a transparent, high-performance, and vectorized AC power flow solver in Python. Key milestones achieved include: (1) rigorous mathematical derivation of the exact analytical polar Jacobian; (2) vectorized NumPy implementation eliminating nested loops; (3) deployment of an interactive Streamlit web dashboard; (4) cross-validation against PyPower matching to 1e-14; and (5) creation of a 34-test automated pytest verification suite.")

    add_heading_with_spacing(doc, "7.2 Directions for Future Research", level=2)
    add_body_p(doc, "Future extensions recommended for Project II and Graduation Thesis include: (1) implementing sparse matrix algorithms using scipy.sparse and KLU; (2) integrating automated generator Q-limit enforcement with anti-windup logic; (3) developing Continuation Power Flow (CPF) to trace P-V bifurcation curves; and (4) coupling the solver with non-linear optimization frameworks for Optimal Power Flow (OPF).")

    doc.add_page_break()

    # ==================== REFERENCES ====================
    add_heading_with_spacing(doc, "References", level=1)
    refs = [
        "[1] J. J. Grainger and W. D. Stevenson, Power System Analysis, McGraw-Hill, New York, 1994.",
        "[2] H. Saadat, Power System Analysis, 3rd ed., PSA Publishing, 2010.",
        "[3] P. Kundur, Power System Stability and Control, McGraw-Hill, New York, 1994.",
        "[4] A. J. Wood, B. F. Wollenberg, and G. B. Sheblé, Power Generation, Operation, and Control, 3rd ed., John Wiley & Sons, Hoboken, NJ, 2013.",
        "[5] W. F. Tinney and C. E. Hart, 'Power flow solution by Newton\\'s method,' IEEE Transactions on Power Apparatus and Systems, vol. PAS-86, no. 11, pp. 1449–1460, Nov. 1967.",
        "[6] B. Stott and O. Alsac, 'Fast decoupled load flow,' IEEE Transactions on Power Apparatus and Systems, vol. PAS-93, no. 3, pp. 859–869, May 1974.",
        "[7] R. D. Zimmerman, C. E. Murillo-Sánchez, and R. J. Thomas, 'MATPOWER: Steady-state operations, planning, and analysis tools for power systems research and education,' IEEE Transactions on Power Systems, vol. 26, no. 1, pp. 12–19, Feb. 2011.",
        "[8] R. Lincoln, 'PYPOWER: A Python port of the MATPOWER electric power system simulation package,' GitHub Repository, 2020.",
        "[9] A. Trias, 'The holomorphic embedding load flow method,' in IEEE Power and Energy Society General Meeting, San Diego, CA, 2012, pp. 1–8.",
        "[10] V. Ajjarapu and C. Christy, 'The continuation power flow: a tool for steady state voltage stability analysis,' IEEE Transactions on Power Systems, vol. 7, no. 1, pp. 416–423, Feb. 1992.",
        "[11] A. Monticelli, State Estimation in Electric Power Systems: A Generalized Approach, Kluwer Academic Publishers, Boston, 1999.",
        "[12] T. Van Cutsem and C. Vournas, Voltage Stability of Electric Power Systems, Springer Science & Business Media, Boston, 2008.",
        "[13] IEEE Standard 399-1997, IEEE Recommended Practice for Industrial and Commercial Power Systems Analysis (Brown Book), IEEE, 1997.",
        "[14] C. R. Harris et al., 'Array programming with NumPy,' Nature, vol. 585, no. 7825, pp. 357–362, Sep. 2020.",
        "[15] W. McKinney, 'Data structures for statistical computing in Python,' in Proceedings of the 9th Python in Science Conference, Austin, TX, 2010, pp. 56–61.",
        "[16] P. Virtanen et al., 'SciPy 1.0: Fundamental algorithms for scientific computing in Python,' Nature Methods, vol. 17, no. 3, pp. 261–272, Mar. 2020.",
        "[17] School of Electrical and Electronic Engineering (SEEE), Hanoi University of Science and Technology, 'Academic Curriculum for Elite Program EE-E18: Power Systems and Renewable Energy,' Hanoi, Vietnam, 2026."
    ]
    for r in refs:
        add_body_p(doc, r, space_after=4)

    doc.add_page_break()

    # ==================== APPENDICES ====================
    add_heading_with_spacing(doc, "Appendix A: Mathematical Proofs of Polar Jacobian Elements", level=1)
    add_heading_with_spacing(doc, "A.1 Rigorous Expansion of Real Power Derivative dP_i / ddelta_i", level=2)
    add_body_p(doc, "Recall the real power injection equation: P_i = |V_i|^2 * G_ii + sum_{k != i} |V_i| * |V_k| * [ G_ik * cos(delta_i - delta_k) + B_ik * sin(delta_i - delta_k) ]. Differentiating with respect to delta_i:")
    add_formula_p(doc, "dP_i / ddelta_i = - sum_{k != i} |V_i| * |V_k| * [ G_ik * sin(delta_ik) - B_ik * cos(delta_ik) ]")
    add_body_p(doc, "From the reactive injection definition: Q_i = -|V_i|^2 * B_ii + sum_{k != i} |V_i| * |V_k| * [ G_ik * sin(delta_ik) - B_ik * cos(delta_ik) ]. Therefore, the summation term equals Q_i + |V_i|^2 * B_ii. Substituting this into the derivative yields:")
    add_formula_p(doc, "dP_i / ddelta_i = -(Q_i + |V_i|^2 * B_ii) = -Q_i - B_ii * |V_i|^2")
    
    add_heading_with_spacing(doc, "A.2 Rigorous Expansion of Real Power Derivative dP_i / d|V_i|", level=2)
    add_body_p(doc, "Differentiating P_i with respect to |V_i|: dP_i / d|V_i| = 2*|V_i|*G_ii + sum_{k != i} |V_k| * [ G_ik * cos(delta_ik) + B_ik * sin(delta_ik) ]. Multiplying and dividing the summation by |V_i|:")
    add_formula_p(doc, "dP_i / d|V_i| = |V_i|*G_ii + (1/|V_i|) * [ |V_i|^2*G_ii + sum_{k != i} |V_i|*|V_k|*(G_ik*cos(delta_ik) + B_ik*sin(delta_ik)) ] = P_i / |V_i| + G_ii * |V_i|")

    doc.add_page_break()

    add_heading_with_spacing(doc, "Appendix B: Tabular Benchmark Case Datasets", level=1)
    add_heading_with_spacing(doc, "B.1 The 3-Bus Educational Benchmark Dataset", level=2)
    b3_bus_data = [
        ("1", "Slack", "0.0", "0.0", "0.0", "0.0", "1.05", "0.0"),
        ("2", "PV", "0.5", "0.0", "0.0", "0.0", "1.02", "0.0"),
        ("3", "PQ", "0.0", "0.0", "0.6", "0.25", "1.00", "0.0")
    ]
    add_styled_table(doc, ["bus", "type", "pg_pu", "qg_pu", "pd_pu", "qd_pu", "vm_pu", "va_deg"], b3_bus_data, col_widths=[0.7, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8])
    
    b3_line_data = [
        ("1", "2", "0.02", "0.06"),
        ("1", "3", "0.08", "0.24"),
        ("2", "3", "0.06", "0.18")
    ]
    add_styled_table(doc, ["from_bus", "to_bus", "r_pu", "x_pu"], b3_line_data, col_widths=[1.5, 1.5, 1.7, 1.7])

    add_heading_with_spacing(doc, "B.2 The 5-Bus Synthetic Educational System Dataset", level=2)
    b5_bus_data = [
        ("1", "Slack", "0.0", "0.0", "0.0", "0.0", "1.05", "0.0"),
        ("2", "PV", "0.9", "0.0", "0.1", "0.04", "1.02", "0.0"),
        ("3", "PQ", "0.0", "0.0", "0.45", "0.18", "1.00", "0.0"),
        ("4", "PQ", "0.0", "0.0", "0.40", "0.15", "1.00", "0.0"),
        ("5", "PQ", "0.0", "0.0", "0.30", "0.10", "1.00", "0.0")
    ]
    add_styled_table(doc, ["bus", "type", "pg_pu", "qg_pu", "pd_pu", "qd_pu", "vm_pu", "va_deg"], b5_bus_data, col_widths=[0.7, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8])
    
    b5_line_data = [
        ("1", "2", "0.02", "0.06"),
        ("1", "3", "0.08", "0.24"),
        ("2", "3", "0.06", "0.18"),
        ("2", "4", "0.04", "0.12"),
        ("3", "4", "0.01", "0.03"),
        ("3", "5", "0.04", "0.12"),
        ("4", "5", "0.03", "0.09")
    ]
    add_styled_table(doc, ["from_bus", "to_bus", "r_pu", "x_pu"], b5_line_data, col_widths=[1.5, 1.5, 1.7, 1.7])

    doc.add_page_break()

    add_heading_with_spacing(doc, "Appendix C: Python Core Solver Source Code", level=1)
    
    src_files = [
        ("src/newton_raphson.py", "C.1 Core Newton–Raphson Solver (`src/newton_raphson.py`)"),
        ("src/jacobian.py", "C.2 Analytical Jacobian Builder (`src/jacobian.py`)"),
        ("src/power_equations.py", "C.3 Power Injections and Mismatches (`src/power_equations.py`)"),
        ("src/ybus.py", "C.4 Nodal Admittance Matrix Builder (`src/ybus.py`)"),
        ("src/branch_flows.py", "C.5 Branch Flows and Loss Engine (`src/branch_flows.py`)"),
        ("src/validation.py", "C.6 Defensive Input Validation Engine (`src/validation.py`)"),
        ("src/data_models.py", "C.7 Data Models and CSV Ingestion (`src/data_models.py`)")
    ]
    for rel_path, heading in src_files:
        full_p = os.path.join(WORKSPACE, rel_path)
        if os.path.exists(full_p):
            add_heading_with_spacing(doc, heading, level=2)
            with open(full_p, "r", encoding="utf-8") as f:
                code_content = f.read()
            add_code_block(doc, code_content)

    doc.add_page_break()

    add_heading_with_spacing(doc, "Appendix D: Streamlit Web Application Code", level=1)
    add_heading_with_spacing(doc, "D.1 Interactive Web Platform (`app.py`)", level=2)
    app_path = os.path.join(WORKSPACE, "app.py")
    if os.path.exists(app_path):
        with open(app_path, "r", encoding="utf-8") as f:
            app_code = f.read()
        add_code_block(doc, app_code)

    print(f"Saving document to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print("Document successfully created and saved!")

if __name__ == "__main__":
    generate_full_word_report()
